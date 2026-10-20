"""Builds small PDF, DOCX and PPTX files for the document tests, so no binary fixtures are kept."""
import textwrap
from pathlib import Path
from typing import List, Sequence

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

BODY_SIZE = 11
LEFT = 72
BOTTOM = 760


class PdfBuilder:
    """Writes text-based PDFs with headings, paragraphs, tables and subscripts."""

    def __init__(self) -> None:
        self.doc = pymupdf.open()
        self.page = None
        self.y = 0.0
        self.new_page()

    def new_page(self) -> "PdfBuilder":
        self.page = self.doc.new_page()
        self.y = 72.0
        return self

    def _room(self, height: float) -> None:
        if self.y + height > BOTTOM:
            self.new_page()

    def heading(self, text: str, size: int = 16) -> "PdfBuilder":
        self._room(size * 2.5)
        self.page.insert_text((LEFT, self.y + size), text, fontsize=size, fontname="hebo")
        self.y += size * 2
        return self

    def line(self, text: str, size: int = BODY_SIZE) -> "PdfBuilder":
        self._room(size * 1.6)
        self.page.insert_text((LEFT, self.y + size), text, fontsize=size, fontname="helv")
        self.y += size * 1.6
        return self

    def paragraph(self, text: str, size: int = BODY_SIZE) -> "PdfBuilder":
        for wrapped in textwrap.wrap(text, width=95):
            self.line(wrapped, size)
        self.y += size * 0.8
        return self

    def subscript_line(self, before: str, sub: str, after: str) -> "PdfBuilder":
        """'before' + a lowered, smaller 'sub' + 'after', as real PDFs set H2O."""
        self._room(30)
        base = self.y + BODY_SIZE
        x = LEFT
        for text, size, drop in ((before, BODY_SIZE, 0), (sub, 7, 3), (after, BODY_SIZE, 0)):
            self.page.insert_text((x, base + drop), text, fontsize=size, fontname="helv")
            x += pymupdf.get_text_length(text, fontname="helv", fontsize=size)
        self.y += BODY_SIZE * 2.2
        return self

    def table(self, rows: Sequence[Sequence[str]], column_width: int = 100) -> "PdfBuilder":
        row_height = 20
        self._room(row_height * len(rows) + 20)
        top = self.y
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                self.page.insert_text(
                    (LEFT + c * column_width + 5, top + r * row_height + 14),
                    cell,
                    fontsize=10,
                    fontname="helv",
                )
        width = column_width * len(rows[0])
        for r in range(len(rows) + 1):
            self.page.draw_line((LEFT, top + r * row_height), (LEFT + width, top + r * row_height))
        for c in range(len(rows[0]) + 1):
            self.page.draw_line(
                (LEFT + c * column_width, top),
                (LEFT + c * column_width, top + row_height * len(rows)),
            )
        self.y = top + row_height * len(rows) + 20
        return self

    def image_page(self) -> "PdfBuilder":
        """A page holding only a picture, like a scanned page."""
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
        pixmap.set_rect(pixmap.irect, (200, 10, 10))
        self.page.insert_image(pymupdf.Rect(72, 72, 272, 172), pixmap=pixmap)
        self.y = BOTTOM  # the next write starts a new page
        return self

    def save(self, path: Path) -> str:
        self.doc.save(str(path))
        self.doc.close()
        return str(path)


def sentences(count: int, prefix: str = "Sentence") -> List[str]:
    """Distinct, numbered sentences, so tests can tell exactly which text a chunk holds."""
    return [f"{prefix} number {i:04d} explains one idea about hash tables." for i in range(count)]


def build_docx(path: Path) -> str:
    document = Document()
    document.add_heading("Course Notes", level=1)
    document.add_paragraph("Intro paragraph before the table.")
    document.add_heading("Costs", level=2)
    table = document.add_table(rows=3, cols=2)
    for r, (left, right) in enumerate([("Structure", "Lookup"), ("Array", "O(n)"), ("Hash", "O(1)")]):
        table.cell(r, 0).text = left
        table.cell(r, 1).text = right
    document.add_paragraph("Paragraph after the table.")
    water = document.add_paragraph("Water is H")
    water.add_run("2").font.subscript = True
    water.add_run("O.")
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("1. A numbered line typed by hand")
    document.save(str(path))
    return str(path)


def build_pptx(path: Path) -> str:
    presentation = Presentation()
    layout_title_and_content = presentation.slide_layouts[1]
    layout_blank = presentation.slide_layouts[6]

    first = presentation.slides.add_slide(layout_title_and_content)
    first.shapes.title.text = "Hashing"
    first.placeholders[1].text_frame.text = "Hash functions map keys to buckets"
    first.placeholders[1].text_frame.add_paragraph().text = "Collisions need a strategy"
    first.notes_slide.notes_text_frame.text = "Mention chaining versus open addressing."

    second = presentation.slides.add_slide(layout_blank)
    box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Free text box on slide two"
    table_shape = second.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    for r, row in enumerate([("Key", "Bucket"), ("apple", "3")]):
        for c, value in enumerate(row):
            table_shape.table.cell(r, c).text = value
    group = second.shapes.add_group_shape()
    grouped = group.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    grouped.text_frame.text = "Text inside a grouped shape"
    presentation.save(str(path))
    return str(path)
